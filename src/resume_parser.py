"""Bounded in-memory parsers. Uploaded bytes are never written to disk."""
import io
import re
import zipfile
from pathlib import PurePath

from docx import Document
from docx.table import Table
from pypdf import PdfReader, apply_configuration

from src.config import MAX_PAGES, MAX_TEXT_CHARS, MAX_UPLOAD_BYTES
from src.models import ParsedDocument

class DocumentError(ValueError):
    """A safe, actionable upload error suitable for display to a user."""


def safe_filename(filename: str) -> str:
    name = filename.replace("\\", "/").rsplit("/", 1)[-1]
    name = re.sub(r"[^\w .()\-]", "_", name).strip(" .")
    suffix = PurePath(name).suffix
    return name[:100 - len(suffix)] + suffix if len(name) > 100 else name or "resume"


def validate_upload(filename: str, data: bytes, allow_txt: bool = False) -> str:
    ext = PurePath(filename.replace("\\", "/")).suffix.lower()
    allowed = {".pdf", ".docx", ".txt"} if allow_txt else {".pdf", ".docx"}
    if ext not in allowed:
        raise DocumentError("Choose a PDF or DOCX document." if not allow_txt else "Choose a PDF, DOCX or UTF-8 TXT document.")
    if not data:
        raise DocumentError("This file is empty. Export the document again and retry.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise DocumentError(f"This file is too large. The limit is {MAX_UPLOAD_BYTES // (1024 * 1024)} MB.")
    if ext == ".pdf" and b"%PDF-" not in data[:1024]:
        raise DocumentError("This file does not contain a valid PDF. Changing an extension does not convert a document.")
    if ext == ".docx" and not data.startswith(b"PK"):
        raise DocumentError("This file does not contain a valid DOCX. Export it as Word .docx and retry.")
    return ext


def _table_text(table: Table) -> list[str]:
    lines: list[str] = []
    for row in table.rows:
        seen = set()
        for cell in row.cells:
            if cell._tc in seen:
                continue
            seen.add(cell._tc)
            lines.extend(p.text for p in cell.paragraphs)
            for nested in cell.tables:
                lines.extend(_table_text(nested))
    return lines


def _docx_text(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        infos = archive.infolist()
        if len(infos) > 500 or sum(i.file_size for i in infos) > 40 * 1024 * 1024:
            raise DocumentError("This Word document expands beyond the safe size limit. Use a simpler export.")
        if "word/document.xml" not in archive.namelist():
            raise DocumentError("This is not a readable Word document.")
        if any("vbaproject" in i.filename.lower() for i in infos):
            raise DocumentError("Documents containing macros are not supported. Export a plain DOCX or PDF.")
        if any(i.filename.endswith(".xml") and i.file_size > 8 * 1024 * 1024 for i in infos):
            raise DocumentError("The document XML is too large. Use a simpler export.")
    doc = Document(io.BytesIO(data))
    lines: list[str] = []
    for section in doc.sections:
        if not section.header.is_linked_to_previous:
            lines.extend(p.text for p in section.header.paragraphs)
            for table in section.header.tables:
                lines.extend(_table_text(table))
    for block in doc.iter_inner_content():
        lines.extend(_table_text(block) if isinstance(block, Table) else [block.text])
    for section in doc.sections:
        if not section.footer.is_linked_to_previous:
            lines.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(lines)


def parse_document(filename: str, data: bytes, *, allow_txt: bool = False) -> ParsedDocument:
    ext = validate_upload(filename, data, allow_txt)
    try:
        return _parse_validated(filename, data, ext)
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("The document could not be read. It may be damaged or use an unsupported format. Export a new PDF/DOCX and retry.") from exc


def _parse_validated(filename: str, data: bytes, ext: str) -> ParsedDocument:
    warnings: list[str] = []
    pages = None
    # Context-local configuration prevents a parser request from changing the
    # limits for another concurrent session.
    with apply_configuration(zlib_maximum_output_length=20 * 1024 * 1024, array_based_stream_maximum_output_length=20 * 1024 * 1024, maximum_declared_stream_length=20 * 1024 * 1024):
        if ext == ".pdf":
            reader = PdfReader(io.BytesIO(data), strict=False)
            if reader.is_encrypted:
                raise DocumentError("This PDF is password-protected. Upload an unlocked copy.")
            pages = len(reader.pages)
            if pages > MAX_PAGES:
                raise DocumentError(f"Use a document with at most {MAX_PAGES} pages.")
            chunks = []
            blank = 0
            for page in reader.pages:
                contents = page.get_contents()
                if contents is not None and len(contents.get_data()) > 5 * 1024 * 1024:
                    raise DocumentError("A PDF page is too complex to process safely. Export a simpler text PDF.")
                chunk = page.extract_text() or ""
                blank += not bool(chunk.strip())
                chunks.append(chunk)
                if sum(map(len, chunks)) > MAX_TEXT_CHARS:
                    raise DocumentError("The extracted document is too long. Upload a shorter resume or job description.")
            text = "\n\n".join(chunks)
            if blank:
                warnings.append(f"{blank} page(s) contained no readable text. Image-only pages need OCR before upload.")
        elif ext == ".docx":
            text = _docx_text(data)
        else:
            text = data.decode("utf-8-sig")
    if len(text) > MAX_TEXT_CHARS:
        raise DocumentError("The extracted document is too long. Upload a shorter document.")
    if len(re.findall(r"\w", text)) < 30 or len(text.split()) < 5:
        raise DocumentError("Too little readable text was found. Scanned PDFs need OCR; export a text-based PDF or DOCX.")
    return ParsedDocument(safe_filename(filename), text, len(data), pages, tuple(warnings))
