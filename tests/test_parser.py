import io
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter
from reportlab.pdfgen import canvas

from src.config import MAX_UPLOAD_BYTES, ROOT
from src.resume_parser import DocumentError, parse_document, safe_filename, validate_upload


def make_pdf(texts):
    stream = io.BytesIO()
    doc = canvas.Canvas(stream)
    for text in texts:
        doc.drawString(40, 750, text)
        doc.showPage()
    doc.save()
    return stream.getvalue()


def test_pdf_extracts_all_pages():
    result = parse_document("two.pdf", make_pdf(["First page has Python SQL and engineering projects", "Second page includes FastAPI Docker and model evaluation"]))
    assert result.pages == 2
    assert "First page" in result.text and "Second page" in result.text


def test_docx_paragraphs_tables_and_header():
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Synthetic resume header"
    doc.add_paragraph("Experience in Python and SQL for data applications")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "FastAPI Docker scikit-learn"
    stream = io.BytesIO()
    doc.save(stream)
    result = parse_document("test.docx", stream.getvalue())
    assert all(term in result.text for term in ["Synthetic", "Python", "FastAPI", "scikit-learn"])


@pytest.mark.parametrize("stem", ["data_analyst", "ml_engineer", "nlp_engineer"])
def test_pdf_docx_samples_have_same_skill_text(stem):
    a = parse_document(stem + ".pdf", (ROOT / "sample_resumes" / (stem + ".pdf")).read_bytes())
    b = parse_document(stem + ".docx", (ROOT / "sample_resumes" / (stem + ".docx")).read_bytes())
    assert " ".join(a.text.split()) == " ".join(b.text.split())


@pytest.mark.parametrize("filename,data", [("run.py", b"print(1)"), ("resume.pdf", b""), ("resume.pdf", b"not a pdf"), ("resume.pdf", b"%PDF-1.7 broken"), ("resume.docx", b"PKbroken"), ("resume.docx", b"bad")])
def test_rejects_bad_files(filename, data):
    with pytest.raises(DocumentError):
        parse_document(filename, data)


def test_size_limit():
    with pytest.raises(DocumentError, match="too large"):
        validate_upload("big.pdf", b"%PDF-" + b"x" * MAX_UPLOAD_BYTES)


def test_scanned_or_empty_pdf():
    with pytest.raises(DocumentError, match="OCR"):
        parse_document("scan.pdf", make_pdf([""]))


def test_partly_empty_pdf_warns():
    result = parse_document("partial.pdf", make_pdf(["", "Python SQL engineering experience includes working on data projects"]))
    assert len(result.warnings) == 1


def test_encrypted_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=600, height=800)
    writer.encrypt("test-password")
    stream = io.BytesIO()
    writer.write(stream)
    with pytest.raises(DocumentError, match="password"):
        parse_document("locked.pdf", stream.getvalue())


def test_too_many_pages():
    with pytest.raises(DocumentError, match="20 pages"):
        parse_document("long.pdf", make_pdf(["Text"] * 21))


def test_zip_expansion_limit():
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * (9 * 1024 * 1024))
    with pytest.raises(DocumentError, match="XML is too large"):
        parse_document("bomb.docx", stream.getvalue())


def test_macros_rejected():
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("word/document.xml", "<doc/>")
        archive.writestr("word/vbaProject.bin", "dummy")
    with pytest.raises(DocumentError, match="macros"):
        parse_document("macro.docx", stream.getvalue())


def test_filename_never_contains_path():
    assert safe_filename("../../private/Resume<script>.pdf") == "Resume_script_.pdf"
    assert safe_filename(r"C:\Users\private\cv.docx") == "cv.docx"


def test_jd_txt_and_invalid_utf8():
    text = b"Required Python SQL Docker and FastAPI skills for API development."
    assert parse_document("job.txt", text, allow_txt=True).text.startswith("Required")
    with pytest.raises(DocumentError):
        parse_document("job.txt", b"\xff" * 80, allow_txt=True)
